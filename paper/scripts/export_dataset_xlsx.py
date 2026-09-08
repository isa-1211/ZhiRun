#!/usr/bin/env python3
"""Export the irrigation model datasets and configuration snapshots to one XLSX."""

from __future__ import annotations

import csv
import gzip
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]
ROOT = PROJECT_ROOT / "灌溉模型" / "灌溉模型"
OUTPUT = Path(__file__).resolve().parents[1] / "data" / "数据集汇总.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def convert_value(value: str):
    """Keep categories as text while converting numeric CSV cells for Excel."""
    if value is None or value == "":
        return None
    value = value.lstrip("\ufeff")
    try:
        if value.isdigit() or (value.startswith("-") and value[1:].isdigit()):
            return int(value)
        number = float(value)
        return number
    except ValueError:
        return value


def style_sheet(ws, widths=None):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    if widths:
        for index, width in widths.items():
            ws.column_dimensions[get_column_letter(index)].width = min(width,  forty := 40)
    else:
        for index, column_cells in enumerate(ws.columns, 1):
            sample = list(column_cells)[:100]
            width = max((len(str(c.value)) if c.value is not None else 0) for c in sample) + 2
            ws.column_dimensions[get_column_letter(index)].width = min(max(width, 10), 28)


def add_csv_sheet(wb, title, path, compressed=False):
    ws = wb.create_sheet(title)
    opener = gzip.open if compressed else open
    with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        for row_index, row in enumerate(reader):
            if row_index == 0:
                ws.append(row)
            else:
                ws.append([convert_value(value) for value in row])
    style_sheet(ws)
    return ws.max_row - 1, ws.max_column


def flatten(value, prefix=""):
    if isinstance(value, dict):
        items = []
        for key, child in value.items():
            child_prefix = f"{prefix}.{key}" if prefix else str(key)
            items.extend(flatten(child, child_prefix))
        return items
    if isinstance(value, list):
        return [(prefix, json.dumps(value, ensure_ascii=False))]
    return [(prefix, value)]


def add_json_sheet(wb, title, path):
    ws = wb.create_sheet(title)
    ws.append(["字段路径", "值"])
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    for key, value in flatten(data):
        ws.append([key, value])
    style_sheet(ws, {1: 42, 2: 80})
    return ws.max_row - 1, 2


def main():
    wb = Workbook()
    overview = wb.active
    overview.title = "数据说明"
    overview.append(["项目", "内容"])
    overview_rows = [
        ("导出时间", datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %z")),
        ("数据根目录", str(ROOT)),
        ("策略样本来源", "data/processed/policy_v2_samples.csv.gz"),
        ("天气数据来源", "data/weather/hohhot_nasa_power_daily.csv"),
        ("说明", "CSV/GZ 数据保留全部原始字段；JSON 配置按字段路径展开。"),
    ]
    for row in overview_rows:
        overview.append(row)
    style_sheet(overview, {1: 18, 2: 80})

    stats = []
    stats.append(("策略样本", *add_csv_sheet(
        wb, "策略样本", ROOT / "data" / "processed" / "policy_v2_samples.csv.gz", compressed=True
    )))
    stats.append(("天气日数据", *add_csv_sheet(
        wb, "天气日数据", ROOT / "data" / "weather" / "hohhot_nasa_power_daily.csv"
    )))
    for title, relative in [
        ("环境快照", Path("data/environment/field_40.72000_111.55000.json")),
        ("作物配置", Path("configs/crops.json")),
        ("土壤配置", Path("configs/soil_profiles.json")),
        ("硬件配置", Path("configs/hardware.json")),
        ("传感器示例", Path("configs/sensor_snapshot.example.json")),
        ("天气来源", Path("data/weather/source.json")),
        ("模型指标", Path("models/policy_v2_metrics.json")),
    ]:
        stats.append((title, *add_json_sheet(wb, title, ROOT / relative)))

    overview.append([])
    overview.append(["工作表", "数据行数", "列数"])
    for item in stats:
        overview.append(item)
    for cell in overview[overview.max_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    # Remove any accidental formula interpretation in free-form text cells.
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + cell.value

    wb.save(OUTPUT)
    print(f"写入: {OUTPUT}")
    for name, rows, cols in stats:
        print(f"{name}: {rows} 行 x {cols} 列")


if __name__ == "__main__":
    main()
